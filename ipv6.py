import os
import re
import io
import argparse
from google.oauth2.credentials import Credentials
from google.auth.transport.requests import Request
from google_auth_oauthlib.flow import InstalledAppFlow
from googleapiclient.discovery import build
from googleapiclient.errors import HttpError
from googleapiclient.http import MediaIoBaseDownload
import docx

# --- CONFIGURAÇÕES GLOBAIS (COPIADAS DO app.py) ---
SCOPES_GOOGLE = ['https://www.googleapis.com/auth/drive', 'https://www.googleapis.com/auth/spreadsheets','https://www.googleapis.com/auth/documents', 'openid', 'https://www.googleapis.com/auth/userinfo.email', 'https://www.googleapis.com/auth/userinfo.profile']
CREDENTIALS_PATH = os.path.join(os.path.dirname(__file__), 'credentials_ipv6.json')
TOKEN_PATH = os.path.join(os.path.dirname(__file__), 'token.json')

def get_credentials():
    creds = None
    if os.path.exists(TOKEN_PATH):
        creds = Credentials.from_authorized_user_file(TOKEN_PATH, SCOPES_GOOGLE)
    if not creds or not creds.valid:
        if creds and creds.expired and creds.refresh_token:
            creds.refresh(Request())
        else:
            flow = InstalledAppFlow.from_client_secrets_file(
                CREDENTIALS_PATH, SCOPES_GOOGLE)
            creds = flow.run_local_server(port=8080)
        with open(TOKEN_PATH, 'w') as token:
            token.write(creds.to_json())
    return creds

def obter_ipv6_do_drive(drive_service, id_pasta_cliente):
    """
    Busca por um arquivo .docx na pasta 'Infraestrutura' de um cliente no Google Drive,
    extrai um endereço IPv6 e o retorna.
    """
    try:
        print(f"DEBUG: Iniciando busca de IPv6 na pasta do cliente com ID: {id_pasta_cliente}")
        # 1. Encontrar a subpasta 'Infraestrutura'
        query_infra = f"name = 'Infraestrutura' and mimeType = 'application/vnd.google-apps.folder' and '{id_pasta_cliente}' in parents and trashed = false"
        results_infra = drive_service.files().list(q=query_infra, fields="files(id, name)").execute()
        items_infra = results_infra.get('files', [])
        if not items_infra:
            print(f"DEBUG: Pasta 'Infraestrutura' não encontrada na pasta do cliente com ID '{id_pasta_cliente}'.")
            return None
        id_pasta_infra = items_infra[0]['id']
        print(f"DEBUG: Pasta 'Infraestrutura' encontrada (ID: {id_pasta_infra})")

        # 2. Listar arquivos .docx na pasta 'Infraestrutura'
        query_docx = f"mimeType = 'application/vnd.openxmlformats-officedocument.wordprocessingml.document' and '{id_pasta_infra}' in parents and trashed = false"
        results_docx = drive_service.files().list(q=query_docx, fields="files(id, name)").execute()
        files_docx = results_docx.get('files', [])
        print(f"DEBUG: Arquivos .docx encontrados na pasta 'Infraestrutura': {[f['name'] for f in files_docx]}")
        
        # 3. Filtrar o arquivo correto (que não contém 'broker')
        target_file = None
        for file in files_docx:
            if 'broker' not in file['name'].lower():
                target_file = file
                break
        
        if not target_file:
            print("DEBUG: Nenhum arquivo .docx válido (sem 'broker' no nome) encontrado.")
            return None
        
        print(f"DEBUG: Arquivo .docx alvo selecionado: {target_file['name']} (ID: {target_file['id']})")

        # 4. Baixar e ler o conteúdo do arquivo .docx
        request_download = drive_service.files().get_media(fileId=target_file['id'])
        fh = io.BytesIO()
        downloader = MediaIoBaseDownload(fh, request_download)
        done = False
        while done is False:
            status, done = downloader.next_chunk()
        
        fh.seek(0)
        document = docx.Document(fh)
        full_text = []
        for para in document.paragraphs:
            full_text.append(para.text)
        
        content = '\n'.join(full_text)

        # 5. Extrair o IPv6 de parágrafos
        for para in document.paragraphs:
            print(f"DEBUG: Verificando parágrafo: {para.text}")
            if 'vpn animati' in para.text.lower():
                print(f"DEBUG: 'vpn animati' encontrado no parágrafo: {para.text}")
                match = re.search(r'(fd[0-9a-fA-F:]+)', para.text)
                if match:
                    ipv6 = match.group(1)
                    print(f"DEBUG: IPv6 encontrado: {ipv6}")
                    return ipv6

        # 6. Extrair o IPv6 de tabelas
        for table in document.tables:
            for row in table.rows:
                for cell in row.cells:
                    for para in cell.paragraphs:
                        print(f"DEBUG: Verificando parágrafo na célula: {para.text}")
                        if 'vpn animati' in para.text.lower():
                            print(f"DEBUG: 'vpn animati' encontrado na célula: {para.text}")
                            match = re.search(r'(fd[0-9a-fA-F:]+)', para.text)
                            if match:
                                ipv6 = match.group(1)
                                print(f"DEBUG: IPv6 encontrado: {ipv6}")
                                return ipv6
        
        print("DEBUG: IPv6 não encontrado no documento.")
        return None

    except HttpError as error:
        print(f"ERRO: Ocorreu um erro na API do Google Drive: {error}")
        return None
    except Exception as e:
        print(f"ERRO: Ocorreu um erro inesperado em obter_ipv6_do_drive: {e}")
        return None

def main():
    parser = argparse.ArgumentParser(description='Busca IPv6 em um arquivo .docx no Google Drive.')
    parser.add_argument('folder_url', type=str, help='A URL da pasta do cliente no Google Drive.')
    args = parser.parse_args()

    creds = get_credentials()
    drive_service = build('drive', 'v3', credentials=creds)
    
    folder_id_match = re.search(r'/folders/([a-zA-Z0-9_-]+)', args.folder_url)
    if not folder_id_match:
        print("ERRO: URL da pasta inválida. Não foi possível extrair o ID da pasta.")
        return
        
    folder_id = folder_id_match.group(1)
    
    ipv6 = obter_ipv6_do_drive(drive_service, folder_id)
    
    if ipv6:
        print(f"IPv6 encontrado: {ipv6}")
    else:
        print("Não foi possível encontrar o IPv6.")

if __name__ == '__main__':
    main()